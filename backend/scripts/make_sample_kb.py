"""Generate the sample knowledge base files.

Run once:  python scripts/make_sample_kb.py

Creates in knowledge_base/:
  05_ml_course_details.pdf                       (single-topic PDF)
  06_fees_scholarships.docx                      (Word document)
  Greenfield_University_Complete_Handbook.pdf    (the big multi-page PDF handbook)
"""
import sys
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parent.parent
KB_DIR = BACKEND_DIR / "knowledge_base"

ML_COURSE_TEXT = """Machine Learning (CSE-4021) - Course Details

Machine Learning (course code CSE-4021) is a 3-credit elective course offered to final-year B.Sc. CSE students in Semester 7. The course instructor is Dr. Elena Rodriguez, Professor of Computer Science, and the teaching assistant is Mr. Daniel Kim.

Classes are held on Sunday and Tuesday from 10:00 AM to 11:30 AM in Room 301 of the Technology Building. The weekly machine learning lab session takes place on Tuesday from 2:00 PM to 5:00 PM in the Artificial Intelligence Lab.

The course covers the following major topics: introduction to machine learning and Python tooling, linear and logistic regression, decision trees and random forests, support vector machines, unsupervised learning and clustering, principal component analysis, neural networks, convolutional neural networks for images, recurrent neural networks for sequences, an introduction to transformers, and model evaluation with cross-validation and standard metrics.

The required textbook is Pattern Recognition and Machine Learning by Christopher Bishop. Two recommended references are Hands-On Machine Learning with Scikit-Learn, Keras and TensorFlow by Aurelien Geron, and Deep Learning by Goodfellow, Bengio and Courville.

Grading is based on a midterm examination worth 30 percent, a semester-long machine learning project worth 40 percent, and a final examination worth 30 percent. Students must submit the project demo before the 13th week. Attendance below 70 percent disqualifies a student from sitting the final examination.

The prerequisite for this course is Artificial Intelligence (CSE-3201). Students are expected to be comfortable with Python programming and linear algebra. The course uses Python with NumPy, pandas, scikit-learn and PyTorch in the lab sessions.
"""

FEES_TEXT = """Tuition Fees and Scholarships

Tuition fees at Greenfield University are charged per semester. The tuition fee for the Computer Science and Engineering program is 4,200 dollars per semester. The tuition fee for Electrical and Electronic Engineering is 4,000 dollars per semester. The tuition fee for Business Administration is 3,200 dollars per semester. The tuition fee for Medicine is 6,500 dollars per semester.

All students pay a one-time admission fee of 500 dollars and a semester registration fee of 150 dollars. The hostel accommodation fee is 800 dollars per semester, which includes meals from the hostel mess.

Merit scholarships are awarded every semester. The Chancellor Scholarship grants a 100 percent tuition waiver to students with a CGPA of 3.90 or above. The Dean's List Scholarship grants a 25 percent tuition waiver to students with a CGPA between 3.70 and 3.89. The Vice-Chancellor Scholarship grants a 50 percent tuition waiver to the top two students of each faculty every year.

Need-based financial aid covers up to 40 percent of tuition for students from low-income families; applications are due in the first week of each semester at the Student Affairs Office in Room 108 of the Administration Building.

Sibling discounts give a 15 percent tuition waiver to the younger sibling of an enrolled student. Full-time university employees and their children receive a 50 percent waiver.

Scholarship results are published within three weeks of the semester's second week on the student portal. Scholarships are revoked if a student's CGPA falls below the eligibility threshold in two consecutive semesters.

For questions about fees, contact the Accounts Office at accounts@greenfield.edu or +1-555-0140. For scholarship questions, contact the Scholarship Desk at scholarships@greenfield.edu.
"""

# (heading, source file in knowledge_base/) for the big handbook
HANDBOOK_SECTIONS = [
    ("1. University Overview", "01_university_overview.txt"),
    ("2. Admissions", "02_admissions.md"),
    ("3. Department of Computer Science and Engineering", "03_cse_department.txt"),
    ("4. B.Sc. CSE Syllabus Outline", "04_btech_cse_syllabus.md"),
    ("5. Machine Learning Course (CSE-4021)", None),          # from ML_COURSE_TEXT
    ("6. Tuition Fees and Scholarships", None),               # from FEES_TEXT
    ("7. Central Library Services", "07_library_services.txt"),
    ("8. Hostel and Accommodation Facilities", "08_hostel_facilities.md"),
    ("9. Sports, Clubs and Student Activities", "09_sports_clubs.txt"),
    ("10. Contact Directory", "10_contact_directory.md"),
    ("11. Academic Calendar and Examinations", "11_exam_calendar.txt"),
    ("12. Campus News", "12_campus_news.html"),
]


def _read_section(heading: str, source: str | None) -> tuple[str, str]:
    if source is None:
        text = ML_COURSE_TEXT if "Machine Learning" in heading else FEES_TEXT
        return heading, text
    path = KB_DIR / source
    raw = path.read_text(encoding="utf-8")
    if source.endswith(".html"):
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style", "nav", "header", "footer"]):
            tag.decompose()
        raw = soup.get_text("\n", strip=True)
    raw = raw.replace("#", "").replace("|", " ").replace("- ", "* ")
    lines = [line.strip() for line in raw.splitlines()]
    return heading, "\n".join(line for line in lines if line)


def make_ml_pdf() -> Path:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", "B", 16)
    pdf.multi_cell(0, 10, "Machine Learning (CSE-4021) - Course Details", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(4)
    pdf.set_font("helvetica", "", 11)
    for paragraph in ML_COURSE_TEXT.split("\n\n"):
        pdf.multi_cell(0, 6, paragraph.strip(), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(2)
    out = KB_DIR / "05_ml_course_details.pdf"
    pdf.output(str(out))
    return out


def make_fees_docx() -> Path:
    import docx

    document = docx.Document()
    document.add_heading("Tuition Fees and Scholarships", level=1)
    for paragraph in FEES_TEXT.split("\n\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())
    out = KB_DIR / "06_fees_scholarships.docx"
    document.save(str(out))
    return out


def make_handbook_pdf() -> Path:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)

    # ---- cover page ----
    pdf.add_page()
    pdf.ln(60)
    pdf.set_font("helvetica", "B", 28)
    pdf.multi_cell(0, 14, "Greenfield University", new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")
    pdf.ln(4)
    pdf.set_font("helvetica", "", 16)
    pdf.multi_cell(0, 10, "Complete Student Handbook", new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")
    pdf.ln(6)
    pdf.set_font("helvetica", "I", 12)
    pdf.multi_cell(0, 8, "Knowledge for a Better Tomorrow", new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")
    pdf.ln(10)
    pdf.set_font("helvetica", "", 11)
    pdf.multi_cell(0, 8, "2026-27 Edition  |  www.greenfield.edu  |  +1-555-0100", new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")

    # ---- table of contents ----
    pdf.add_page()
    pdf.set_font("helvetica", "B", 16)
    pdf.multi_cell(0, 10, "Table of Contents", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(2)
    pdf.set_font("helvetica", "", 12)
    for heading, _ in HANDBOOK_SECTIONS:
        pdf.multi_cell(0, 8, heading, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(6)
    pdf.set_font("helvetica", "I", 10)
    pdf.multi_cell(0, 7,
        "This handbook is the official knowledge base of the Greenfield University "
        "information service. It is intended for students, applicants, and staff.")

    # ---- sections ----
    for heading, source in HANDBOOK_SECTIONS:
        title, body = _read_section(heading, source)
        pdf.add_page()
        pdf.set_font("helvetica", "B", 16)
        pdf.multi_cell(0, 10, title, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(2)
        pdf.set_font("helvetica", "", 11)
        for paragraph in body.split("\n"):
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            pdf.multi_cell(0, 6, paragraph, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(1)

    out = KB_DIR / "Greenfield_University_Complete_Handbook.pdf"
    pdf.output(str(out))
    return out


def main() -> None:
    KB_DIR.mkdir(parents=True, exist_ok=True)
    for maker in (make_ml_pdf, make_fees_docx, make_handbook_pdf):
        path = maker()
        print(f"Created {path.name} ({path.stat().st_size} bytes)")


if __name__ == "__main__":
    sys.exit(main())
